#!/usr/bin/env python3
"""Convert manuscript_v2.md (markdown) to a Word file with python-docx.

Usage (env PlantCad2, python-docx 1.2.0 installed 2026-10-12):
    python build_docx.py                      # manuscript_v2.md -> manuscript_v2.docx
    python build_docx.py -i manuscript_v2.md -o manuscript_v2_draft.docx

Handles: headings (#, ##, ###, ####), paragraphs with inline **bold**, *italic*, `code`,
bullet lists ('- '), numbered-looking reference lines ('12. …' kept as plain paragraphs),
markdown tables (| … |) -> Word tables (Table Grid, 8 pt), figures: before a paragraph that
starts with '**Figure N.**' the PNG from paper/figures/ (FIGURES dict) is inserted at 16 cm width.
Font: Palatino Linotype 10 pt (MDPI-like), A4, 2.5 cm margins. Citations [n] stay as text.
"""
import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
PAPER = HERE.parent
FIGURES = {
    "1": PAPER / "figures" / "Fig1_workflow.png",
    "2": PAPER / "figures" / "Fig2_screen.png",
    "3": PAPER / "figures" / "Fig3_trees.png",
    "4": PAPER / "figures" / "Fig4_islands.png",
    "5": PAPER / "figures" / "Fig5_boletales_genera.png",
}
FONT = "Palatino Linotype"
SIZE = 10
FIG_WIDTH_CM = 16

INLINE = re.compile(r"(\*\*.+?\*\*|\*(?!\s).+?(?<!\s)\*|`[^`]+`)")


def set_font(run, bold=None, italic=None, size=None, mono=False):
    run.font.name = "Consolas" if mono else FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), "Consolas" if mono else FONT)
    run.font.size = Pt(size or SIZE)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(par, text, size=None, bold_base=False):
    """Add runs with **bold** / *italic* / `code` parsed from markdown text."""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            set_font(par.add_run(text[pos:m.start()]), bold=bold_base, size=size)
        tok = m.group(0)
        if tok.startswith("**"):
            inner = tok[2:-2]
            # italic inside bold, e.g. **Figure 2.** is plain; **Table 3. *x*** rare — handle *…*
            sub = 0
            for mm in re.finditer(r"\*(?!\s).+?(?<!\s)\*", inner):
                if mm.start() > sub:
                    set_font(par.add_run(inner[sub:mm.start()]), bold=True, size=size)
                set_font(par.add_run(mm.group(0)[1:-1]), bold=True, italic=True, size=size)
                sub = mm.end()
            if sub < len(inner):
                set_font(par.add_run(inner[sub:]), bold=True, size=size)
        elif tok.startswith("`"):
            set_font(par.add_run(tok[1:-1]), size=size, mono=True)
        else:
            inner = tok[1:-1]
            sub = 0
            for mm in re.finditer(r"\*\*.+?\*\*", inner):
                if mm.start() > sub:
                    set_font(par.add_run(inner[sub:mm.start()]), italic=True, bold=bold_base, size=size)
                set_font(par.add_run(mm.group(0)[2:-2]), italic=True, bold=True, size=size)
                sub = mm.end()
            if sub < len(inner):
                set_font(par.add_run(inner[sub:]), italic=True, bold=bold_base, size=size)
        pos = m.end()
    if pos < len(text):
        set_font(par.add_run(text[pos:]), bold=bold_base, size=size)


def style_doc(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(SIZE)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.0
    for lvl, size in ((1, 16), (2, 13), (3, 11), (4, 10)):
        h = doc.styles[f"Heading {lvl}"]
        h.font.name = FONT
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.italic = lvl == 4
        h.font.color.rgb = None
        h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.5))


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c or "---") for c in r)]
    if not cells:
        return
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # column widths for Table 1 (2026-10-12aa): the two host columns get most of the 16 cm text width
    widths = None
    if any("rhizobial T-DNA" in c for c in cells[0]) and ncol == 6:
        widths = [1.4, 1.9, 1.9, 1.9, 4.5, 4.4]
    if widths:
        t.autofit = False
        for j, wcm in enumerate(widths):
            for i in range(len(cells)):
                t.cell(i, j).width = Cm(wcm)
    for i, r in enumerate(cells):
        for j in range(ncol):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, r[j] if j < len(r) else "", size=8, bold_base=(i == 0))
    doc.add_paragraph()


def add_figure(doc, path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(FIG_WIDTH_CM))


def build(md_path: Path, out: Path):
    doc = Document()
    style_doc(doc)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i, n = 0, len(lines)
    n_fig = n_tab = n_par = 0
    missing = []
    while i < n:
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            lvl = len(m.group(1))
            h = doc.add_heading("", level=lvl)
            add_inline(h, m.group(2), size={1: 16, 2: 13, 3: 11, 4: 10}[lvl], bold_base=True)
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < n and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            n_tab += 1
            continue
        if line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
            n_par += 1
            i += 1
            continue
        fm = re.match(r"^\*\*Figure (\d+)\.\*\*", line)
        if fm:
            path = FIGURES.get(fm.group(1))
            if path and path.exists():
                add_figure(doc, path)
                n_fig += 1
            else:
                missing.append(fm.group(1))
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, line)
        n_par += 1
        i += 1
    # highlight open marks for the co-author (2026-10-13): [CHECK…], [PENDING…], [REF], [TO FILL…], [VERIFY WITH AUTHOR…]
    from docx.enum.text import WD_COLOR_INDEX
    MARK = re.compile(r"\[(CHECK|PENDING|REF|TO FILL|VERIFY WITH AUTHOR)")
    def _paras(d):
        for p in d.paragraphs: yield p
        for t in d.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs: yield p
    for p in _paras(doc):
        if MARK.search(p.text):
            for r in p.runs:
                if MARK.search(r.text) or (r.text.strip() and MARK.search(p.text) and len(p.runs) == 1):
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.save(out)
    return n_par, n_tab, n_fig, missing


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("-i", "--inp", default=str(HERE / "manuscript_v2.md"))
    ap.add_argument("-o", "--out", default=str(HERE / "manuscript_v2.docx"))
    a = ap.parse_args()
    n_par, n_tab, n_fig, missing = build(Path(a.inp), Path(a.out))
    print(f"written {a.out}: paragraphs {n_par}, tables {n_tab}, figures {n_fig}, "
          f"missing figure files: {missing or 'none'}")
    # verify by re-opening
    d = Document(a.out)
    print(f"re-opened: paragraphs {len(d.paragraphs)}, tables {len(d.tables)}, "
          f"inline shapes {len(d.inline_shapes)}, sections {len(d.sections)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
