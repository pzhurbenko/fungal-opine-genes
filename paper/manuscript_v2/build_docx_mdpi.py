#!/usr/bin/env python3
"""Build the submission DOCX on the official MDPI Biology Word template (2026-10-13).

Usage (env PlantCad2):  python build_docx_mdpi.py [-i manuscript_v2_clean.md] [-o manuscript_v2_MDPI.docx]
Template: paper/submission/mdpi_template/biology-template.docx (the MDPI .dotx with its content type rewritten;
downloaded from the Wayback Machine copy of https://www.mdpi.com/files/word-templates/biology-template.dot).
The template body is cleared (headers/footers/section settings kept) and the manuscript is written with the
template's paragraph styles: MDPI_1.x title block, MDPI_2.x headings, MDPI_3.1_text body, MDPI_3.8_bullet,
MDPI_4.x tables, MDPI_5.x figures, MDPI_6.2_BackMatter, MDPI_7.1_References. Marks [TO FILL …] / [VERIFY WITH AUTHOR …]
are highlighted in yellow.
"""
import argparse, re
from pathlib import Path
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
PAPER = HERE.parent
TEMPLATE = PAPER / "submission" / "mdpi_template" / "biology-template.docx"
FIGURES = {n: PAPER / "figures" / f for n, f in (("1", "Fig1_workflow.png"), ("2", "Fig2_screen.png"), ("3", "Fig3_trees.png"),
                                                 ("4", "Fig4_islands.png"), ("5", "Fig5_boletales_genera.png"))}
INLINE = re.compile(r"\*\*(?!\s).+?(?<!\s)\*\*|\*(?!\s)[^*]+?(?<!\s)\*|`[^`]+`")
MARK = re.compile(r"\[(CHECK|PENDING|REF|TO FILL|VERIFY WITH AUTHOR)")
BACK = ["Supplementary Materials", "Author Contributions", "Funding", "Institutional Review Board Statement",
        "Informed Consent Statement", "Data Availability Statement", "Acknowledgments", "Conflicts of Interest"]

def runs(par, text, bold=False, italic=False):
    """markdown inline (**bold**, *italic*, `code`) -> runs; nested *…* inside **…** handled."""
    pos = 0
    def add(t, b, i):
        r = par.add_run(t); r.bold = b or None; r.italic = i or None
        if MARK.search(t): r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    for m in INLINE.finditer(text):
        if m.start() > pos: add(text[pos:m.start()], bold, italic)
        tok = m.group(0)
        if tok.startswith("**"):
            inner = tok[2:-2]; sub = 0
            for mm in re.finditer(r"\*(?!\s).+?(?<!\s)\*", inner):
                if mm.start() > sub: add(inner[sub:mm.start()], True, italic)
                add(mm.group(0)[1:-1], True, True); sub = mm.end()
            if sub < len(inner): add(inner[sub:], True, italic)
        elif tok.startswith("`"):
            add(tok[1:-1], bold, italic)
        else:
            add(tok[1:-1], bold, True)
        pos = m.end()
    if pos < len(text): add(text[pos:], bold, italic)

def para(doc, style, text="", **kw):
    p = doc.add_paragraph(style=S(doc, style))
    if text: runs(p, text, **kw)
    return p

ALIASES = {"MDPI_6.2_BackMatter": ["MDPI_6.2_back_matter"], "MDPI_6.3_Notes": ["MDPI_6.3_notes"],
           "MDPI_7.1_References": ["MDPI_8.1_references"]}
def S(doc, name):
    """Resolve a style name across template versions (2024: MDPI_6.2_BackMatter/MDPI_7.1_References; 2026: MDPI_6.2_back_matter/MDPI_8.1_references)."""
    names = {st.name for st in doc.styles}
    for cand in [name] + ALIASES.get(name, []):
        if cand in names: return cand
    raise KeyError(name)

def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if not child.tag.endswith("}sectPr"): body.remove(child)

def style_numbered(doc, name):
    st = doc.styles[name]
    return st.element.find(qn("w:pPr")) is not None and st.element.find(qn("w:pPr")).find(qn("w:numPr")) is not None

def add_table(doc, block):
    rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in block if not re.match(r"^\|\s*:?-{2,}", r)]
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol); t.style = doc.styles["Table Grid"]
    for i, r in enumerate(rows):
        for j in range(ncol):
            cell = t.cell(i, j); cell.text = ""
            p = cell.paragraphs[0]; p.style = doc.styles[S(doc, "MDPI_4.2_table_body")]
            runs(p, r[j] if j < len(r) else "", bold=(i == 0))
    doc.add_paragraph(style=S(doc, "MDPI_3.2_text_no_indent"))
    return t

def build(md: Path, out: Path):
    doc = Document(str(TEMPLATE)); clear_body(doc)
    lines = md.read_text(encoding="utf-8").splitlines()
    refs_numbered = style_numbered(doc, S(doc, "MDPI_7.1_References"))
    i, n = 0, len(lines); section = None; title_done = False; n_fig = n_tab = 0
    para(doc, "MDPI_1.1_article_type", "Article")
    while i < n:
        line = lines[i]
        if not line.strip() or re.match(r"^-{3,}\s*$", line): i += 1; continue
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            lvl, text = len(m.group(1)), m.group(2).strip()
            if lvl == 1 and not title_done:
                para(doc, "MDPI_1.2_title", text); title_done = True
                # author block: following non-empty lines until a blank line
                i += 1; blk = []
                while i < n and not lines[i].strip(): i += 1
                while i < n and lines[i].strip() and not lines[i].startswith("#"): blk.append(lines[i].rstrip("\\ ").strip()); i += 1
                if blk:
                    para(doc, "MDPI_1.3_authornames", blk[0])
                    for b in blk[1:]: para(doc, "MDPI_1.6_affiliation", b)
                continue
            section = text
            if text in ("Simple Summary", "Abstract", "Keywords"):
                i += 1
                while i < n and not lines[i].strip(): i += 1
                body = lines[i].strip() if i < n else ""
                st = "MDPI_1.8_keywords" if text == "Keywords" else "MDPI_1.7_abstract"
                p = para(doc, st); r = p.add_run(f"{text}: "); r.bold = True; runs(p, body)
                i += 1; continue
            if text in BACK:
                i += 1
                while i < n and not lines[i].strip(): i += 1
                body = lines[i].strip() if i < n else ""
                p = para(doc, "MDPI_6.2_BackMatter"); r = p.add_run(f"{text}: "); r.bold = True; runs(p, body)
                i += 1; continue
            if text == "References":
                para(doc, "MDPI_2.1_heading1", "References"); i += 1
                while i < n:
                    l = lines[i].strip()
                    if l:
                        mm = re.match(r"^(\d+)\.\s+(.*)", l)
                        body = mm.group(2) if (mm and refs_numbered) else l
                        para(doc, "MDPI_7.1_References", body)
                    i += 1
                continue
            para(doc, {2: "MDPI_2.1_heading1", 3: "MDPI_2.2_heading2"}.get(lvl, "MDPI_2.3_heading3"), text)
            i += 1; continue
        if line.startswith("|"):
            block = []
            while i < n and lines[i].startswith("|"): block.append(lines[i]); i += 1
            add_table(doc, block); n_tab += 1; continue
        if line.startswith("- ") or line.startswith("* "):
            para(doc, "MDPI_3.8_bullet", line[2:]); i += 1; continue
        fm = re.match(r"^\*\*Figure (\d+)\.\*\*", line)
        tm = re.match(r"^\*\*Table (\d+)\.\*\*", line)
        if fm:
            path = FIGURES.get(fm.group(1))
            if path and path.exists():
                p = doc.add_paragraph(style=S(doc, "MDPI_5.2_figure")); p.add_run().add_picture(str(path), width=Cm(16)); n_fig += 1
            para(doc, "MDPI_5.1_figure_caption", line); i += 1; continue
        if tm:
            para(doc, "MDPI_4.1_table_caption", line); i += 1; continue
        para(doc, "MDPI_3.1_text", line); i += 1
    para(doc, "MDPI_6.3_Notes", "**Disclaimer/Publisher’s Note:** The statements, opinions and data contained in all publications are solely those of the individual author(s) and contributor(s) and not of MDPI and/or the editor(s). MDPI and/or the editor(s) disclaim responsibility for any injury to people or property resulting from any ideas, methods, instructions or products referred to in the content.")
    doc.save(str(out))
    return n_fig, n_tab, refs_numbered

if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("-i", default=str(HERE / "manuscript_v2_clean.md")); ap.add_argument("-o", default=str(HERE / "manuscript_v2_MDPI.docx"))
    a = ap.parse_args(); nf, nt, rn = build(Path(a.i), Path(a.o))
    d = Document(a.o); print(f"written {a.o}: figures {nf}, tables {nt}, paragraphs {len(d.paragraphs)}, references auto-numbered by style: {rn}")
